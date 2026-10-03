import io

import openpyxl
import pytest

from app.core.security import hash_password
from app.models import Hall, MealType, User, UserRole

FP = "dev-1"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
async def accountant_headers(client, db_session):
    u = User(login="buh", password_hash=hash_password("bpass"),
             full_name="Бух", role=UserRole.accountant)
    db_session.add(u)
    await db_session.commit()
    await client.post("/api/v1/auth/login", json={"login": "buh", "password": "bpass"},
                      headers={"X-Device-Fingerprint": FP})
    return {"X-Device-Fingerprint": FP}


async def test_daily_report_export(client, accountant_headers, db_session):
    db_session.add_all([Hall(name="Зал №1"), MealType(name="Мясо", sort_order=1)])
    await db_session.commit()

    r = await client.get("/api/v1/accountant/report/export?date=2026-06-10",
                         headers=accountant_headers)
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX_CT
    assert "attachment" in r.headers["content-disposition"]
    assert "report-2026-06-10.xlsx" in r.headers["content-disposition"]

    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert str(ws.cell(row=1, column=1).value).startswith("Отчёт")


async def test_period_report_export(client, accountant_headers, db_session):
    db_session.add_all([Hall(name="Зал №1"), MealType(name="Мясо", sort_order=1)])
    await db_session.commit()

    r = await client.get(
        "/api/v1/accountant/report/period/export?from=2026-06-01&to=2026-06-03",
        headers=accountant_headers,
    )
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX_CT
    assert "attachment" in r.headers["content-disposition"]
    assert "report-2026-06-01_2026-06-03.xlsx" in r.headers["content-disposition"]

    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert str(ws.cell(row=1, column=1).value).startswith("Отчёт")
