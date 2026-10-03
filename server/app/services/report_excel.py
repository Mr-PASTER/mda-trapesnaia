import io
from collections.abc import Callable
from datetime import date
from typing import cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import MealKind
from app.schemas.report import (
    DailyReportOut, HallPeriodReportOut, HallReportOut, PeriodReportOut,
)

_MEAL_LABELS: dict[MealKind, str] = {
    MealKind.breakfast: "Завтрак",
    MealKind.lunch: "Обед",
    MealKind.snack: "Полдник",
    MealKind.dinner: "Ужин",
}

_TITLE_ROW = 1
_HEADER_ROW = 2
_SUB_HEADER_ROW = 3
_FIRST_DATA_ROW = 4


def _fmt_date(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def _meal_groups(halls: list) -> list[tuple[MealKind, list[str]]]:
    """Column groups are derived from the first hall, whose meals list mirrors
    the active meal types and is identical in shape for every hall."""
    if not halls:
        return []
    groups: list[tuple[MealKind, list[str]]] = []
    for meal in halls[0].meals:
        subcolumns = [t.name for t in meal.by_type]
        subcolumns += [f"Рез. {t.name}" for t in meal.reserve_by_type]
        if subcolumns:
            groups.append((meal.meal_kind, subcolumns))
    return groups


def _write_headers(ws: Worksheet, groups: list[tuple[MealKind, list[str]]]) -> dict:
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    col = 1
    ws.cell(row=_HEADER_ROW, column=col, value="Зал")
    ws.merge_cells(start_row=_HEADER_ROW, start_column=col,
                   end_row=_SUB_HEADER_ROW, end_column=col)
    hall_col = col
    col += 1

    group_layout: list[tuple[MealKind, int, int]] = []
    for meal_kind, subcolumns in groups:
        start = col
        for name in subcolumns:
            ws.cell(row=_SUB_HEADER_ROW, column=col, value=name)
            col += 1
        end = col - 1
        ws.cell(row=_HEADER_ROW, column=start, value=_MEAL_LABELS[meal_kind])
        if end > start:
            ws.merge_cells(start_row=_HEADER_ROW, start_column=start,
                           end_row=_HEADER_ROW, end_column=end)
        group_layout.append((meal_kind, start, end))

    total_col = col
    ws.cell(row=_HEADER_ROW, column=total_col, value="Итого")
    ws.cell(row=_SUB_HEADER_ROW, column=total_col, value="Всего")
    col += 1

    reserve_col = col
    ws.cell(row=_HEADER_ROW, column=reserve_col, value="Итого резерв")
    ws.cell(row=_SUB_HEADER_ROW, column=reserve_col, value="Резерв")
    col += 1

    last_col = col - 1
    for r in (_HEADER_ROW, _SUB_HEADER_ROW):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = bold
            cell.alignment = center

    return {
        "hall_col": hall_col,
        "groups": group_layout,
        "total_col": total_col,
        "reserve_col": reserve_col,
        "last_col": last_col,
    }


def _write_workbook(
    title: str,
    halls: list[HallReportOut] | list[HallPeriodReportOut],
    total_of: Callable,
    reserve_of: Callable,
    grand_total: int,
    grand_reserve_total: int,
) -> bytes:
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Отчёт"

    groups = _meal_groups(halls)
    layout = _write_headers(ws, groups)

    ws.cell(row=_TITLE_ROW, column=1, value=title).font = Font(bold=True, size=13)
    if layout["last_col"] > 1:
        ws.merge_cells(start_row=_TITLE_ROW, start_column=1,
                       end_row=_TITLE_ROW, end_column=layout["last_col"])

    row = _FIRST_DATA_ROW
    for hall in halls:
        meals_by_kind = {meal.meal_kind: meal for meal in hall.meals}
        ws.cell(row=row, column=layout["hall_col"], value=hall.hall_name)
        for meal_kind, start, _end in layout["groups"]:
            meal = meals_by_kind.get(meal_kind)
            if meal is None:
                continue
            values = [t.count for t in meal.by_type]
            values += [t.count for t in meal.reserve_by_type]
            for offset, value in enumerate(values):
                ws.cell(row=row, column=start + offset, value=value)
        ws.cell(row=row, column=layout["total_col"], value=total_of(hall))
        ws.cell(row=row, column=layout["reserve_col"], value=reserve_of(hall))
        row += 1

    bold = Font(bold=True)
    ws.cell(row=row, column=layout["hall_col"], value="ИТОГО").font = bold
    ws.cell(row=row, column=layout["total_col"], value=grand_total).font = bold
    ws.cell(row=row, column=layout["reserve_col"], value=grand_reserve_total).font = bold
    last_row = row

    for c in range(1, layout["last_col"] + 1):
        widest = 0
        for r in range(_HEADER_ROW, last_row + 1):
            value = ws.cell(row=r, column=c).value
            if value is not None:
                widest = max(widest, len(str(value)))
        ws.column_dimensions[get_column_letter(c)].width = min(max(widest + 2, 8), 40)

    ws.freeze_panes = f"A{_FIRST_DATA_ROW}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def daily_report_to_xlsx(report: DailyReportOut) -> bytes:
    return _write_workbook(
        title=f"Отчёт за {_fmt_date(report.date)}",
        halls=report.halls,
        total_of=lambda hall: hall.day_total,
        reserve_of=lambda hall: hall.day_reserve_total,
        grand_total=report.grand_total,
        grand_reserve_total=report.grand_reserve_total,
    )


def period_report_to_xlsx(report: PeriodReportOut) -> bytes:
    return _write_workbook(
        title=(
            f"Отчёт за период {_fmt_date(report.date_from)}–{_fmt_date(report.date_to)}"
        ),
        halls=report.halls,
        total_of=lambda hall: hall.period_total,
        reserve_of=lambda hall: hall.period_reserve_total,
        grand_total=report.grand_total,
        grand_reserve_total=report.grand_reserve_total,
    )
